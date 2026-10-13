"""
Importa a planilha "Levantamento das Máquinas" (.xlsx) para o Parque de Máquinas.

Uso:
    python manage.py import_planilha "caminho/Levantamento das Maquinas.xlsx" [--reset]

Lê cada aba de unidade (Bonsucesso, Taquara, Vilar, Centro, Cabo Frio,
Campo Grande) criando Locais › Setores › Baias, e a aba Transferencias criando
o histórico de movimentações.
"""

from datetime import date, datetime
from pathlib import Path

import openpyxl
from django.core.management.base import BaseCommand, CommandError

from apps.orgs.models import Location
from apps.parque.importer import import_transferencias, import_unidade
from apps.parque.models import Baia, Setor, Transferencia

UNIDADES = ["Bonsucesso", "Taquara", "Vilar", "Centro", "Cabo Frio", "Campo Grande"]


def cell_to_str(v):
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def sheet_rows(ws):
    return [[cell_to_str(c) for c in row] for row in ws.iter_rows(values_only=True)]


class Command(BaseCommand):
    help = "Importa a planilha Levantamento das Máquinas (.xlsx) para o Parque."

    def add_arguments(self, parser):
        parser.add_argument("xlsx_path")
        parser.add_argument(
            "--reset", action="store_true",
            help="Apaga setores/baias/transferências antes de importar.",
        )

    def handle(self, *args, **opts):
        path = Path(opts["xlsx_path"])
        if not path.exists():
            raise CommandError(f"Arquivo não encontrado: {path}")

        wb = openpyxl.load_workbook(path, data_only=True)

        if opts["reset"]:
            Transferencia.objects.all().delete()
            Baia.objects.all().delete()
            Setor.objects.all().delete()
            self.stdout.write(self.style.WARNING("Dados do parque apagados (--reset)."))

        for nome in UNIDADES:
            if nome not in wb.sheetnames:
                self.stdout.write(self.style.WARNING(f"  aba ausente: {nome}"))
                continue
            stats = import_unidade(nome, sheet_rows(wb[nome]))
            self.stdout.write(self.style.SUCCESS(
                f"  {nome}: setores={stats['setores']} baias={stats['baias']}"
            ))

        for sn in wb.sheetnames:
            if sn.lower().startswith("transfer"):
                stats = import_transferencias(sheet_rows(wb[sn]))
                self.stdout.write(self.style.SUCCESS(
                    f"  {sn}: transferências={stats['transferencias']}"
                ))

        self.stdout.write(self.style.SUCCESS(
            f"TOTAIS: Unidades={Location.objects.count()} "
            f"Setores={Setor.objects.count()} Baias={Baia.objects.count()} "
            f"Transferencias={Transferencia.objects.count()}"
        ))
